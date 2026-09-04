"""Inspection report generation — real .docx via python-docx.

``generate_report()`` takes the actual output shapes of Phases 3a/3b/3c
(``OcrResult``, ``CalculationResult``, ``DocumentResult`` — or plain dicts
with the same fields, for callers going through JSON) and renders a Word
document with equipment ID, OCR reading, calculation verdict, timestamp,
and an SOP section.

The SOP section is never silently omitted. Three explicit states, because
silence is ambiguous (it could mean "no SOP needed" or "one was denied"):

* no lookup attempted (``sop_requested`` is None) → "No SOP reference was
  consulted for this inspection."
* lookup attempted and found → cite the resource ID and quote its text.
* lookup attempted but not found → "SOP reference unavailable to this
  user" (deliberately not distinguishing denied vs missing, preserving
  Phase 3c's indistinguishability guarantee all the way into the report).
"""

from __future__ import annotations

import dataclasses
import datetime
import os
import tempfile
from pathlib import Path
from typing import Any, Mapping

from docx import Document

SOP_UNAVAILABLE_TEXT = (
    "SOP reference unavailable to this user: the requested SOP was either "
    "not authorized for this subject or does not exist. The distinction is "
    "withheld so the report cannot be used to probe for restricted documents."
)
SOP_NOT_CONSULTED_TEXT = "No SOP reference was consulted for this inspection."


def _as_dict(value: Any) -> dict:
    if value is None:
        return {}
    if dataclasses.is_dataclass(value):
        return dataclasses.asdict(value)
    if isinstance(value, Mapping):
        return dict(value)
    raise TypeError(f"expected a dataclass or mapping, got {type(value).__name__}")


def _utc_now_iso() -> str:
    return datetime.datetime.now(datetime.timezone.utc).isoformat(timespec="seconds")


def generate_report(inspection_data: dict, output_path: str | None = None) -> Path:
    """Render ``inspection_data`` to a .docx file; return its path.

    Expected keys: ``equipment_id`` (required), ``subject``,
    ``timestamp`` (defaults to current UTC), ``ocr`` (OcrResult/dict),
    ``calculation`` (CalculationResult/dict), ``sop_requested``
    (resource ID looked up, or None), ``sop`` (DocumentResult/dict).
    """
    if not isinstance(inspection_data, dict):
        raise TypeError("inspection_data must be a dict")
    equipment_id = inspection_data.get("equipment_id")
    if not equipment_id or not isinstance(equipment_id, str):
        raise TypeError("inspection_data['equipment_id'] must be a non-empty string")

    subject = inspection_data.get("subject", "?")
    timestamp = inspection_data.get("timestamp") or _utc_now_iso()
    ocr = _as_dict(inspection_data.get("ocr"))
    calc = _as_dict(inspection_data.get("calculation"))
    sop_requested = inspection_data.get("sop_requested")
    sop = _as_dict(inspection_data.get("sop"))

    if calc.get("error"):
        verdict = "ERROR"
        calc_detail = f"Calculation could not run: {calc['error']}"
    elif calc.get("passed"):
        verdict = "PASS"
        calc_detail = (f"Margin: {calc.get('margin_bar')} bar. "
                       f"{calc.get('reason') or ''}".strip())
    else:
        verdict = "FAIL"
        calc_detail = calc.get("reason") or "Failed with no reason given."

    doc = Document()
    doc.add_heading("Inspection Report", level=0)
    doc.add_paragraph(f"Equipment: {equipment_id}")
    doc.add_paragraph(f"Inspected by: {subject}")
    doc.add_paragraph(f"Timestamp (UTC): {timestamp}")

    doc.add_heading("OCR reading", level=1)
    ocr_text = (ocr.get("text") or "").strip()
    doc.add_paragraph(ocr_text if ocr_text else "(no text extracted)")
    conf = ocr.get("mean_confidence")
    if ocr.get("low_confidence"):
        doc.add_paragraph(
            "WARNING: OCR confidence is low "
            f"(mean {conf if conf is not None else '?'}); "
            "readings above should be verified by a human before acting on them."
        )
    elif conf is not None:
        doc.add_paragraph(f"OCR mean confidence: {conf:.1f}")

    doc.add_heading("Calculation", level=1)
    doc.add_paragraph(f"Verdict: {verdict}")
    doc.add_paragraph(calc_detail)

    doc.add_heading("SOP reference", level=1)
    if not sop_requested:
        doc.add_paragraph(SOP_NOT_CONSULTED_TEXT)
    elif sop.get("found") and sop.get("text"):
        doc.add_paragraph(f"Cited SOP: {sop.get('resource_id') or sop_requested}")
        doc.add_paragraph(sop["text"])
    else:
        doc.add_paragraph(f"Requested SOP: {sop_requested}")
        doc.add_paragraph(SOP_UNAVAILABLE_TEXT)

    if output_path is None:
        fd, output_path = tempfile.mkstemp(prefix=f"{equipment_id}-report-",
                                           suffix=".docx")
        os.close(fd)
    target = Path(output_path)
    if target.suffix != ".docx":
        raise ValueError("output_path must end in .docx")
    doc.save(str(target))
    return target
